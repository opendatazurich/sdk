import pandas as pd
import os
import re

from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment


def testexport_json(df: pd.DataFrame, filename="testexport_datasets.json", output_dir: str = "."):
    # 4.1 Testexport for Nils (Initialimport)
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, filename)
    print('write json', path)
    df.to_json(path, orient='records', default_handler=str)


def export_to_excel(df: pd.DataFrame, filename="initialimport.xlsx", output_dir: str = "."):
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, filename)
    print("write excel", path)
    df.to_excel(path, index=False)

def export_to_excel_by_org(df_dataset: pd.DataFrame, df_attribues: pd.DataFrame, filename="initialimport.xlsx", org_col="author_da_gs", output_dir: str = "."):
    """
    Creates an excel file for every unique organization (e.g. dienstabteilung)
    which contains a sheet for datasets and attribues
    
    :param df_dataset: df with dataset info. Needs to have a org_col
    :type df_dataset: pd.DataFrame
    :param df_attribues: df with attribues info. Needs to have a org_col
    :type df_attribues: pd.DataFrame
    :param filename: End of the filename (gets a prefix from org_col)
    :param org_col: Column name the specifies the dienstabteilung
    """
    
    if org_col not in df_dataset.columns:
        raise ValueError(f"Column '{org_col}' missing in df_dataset.")
    if org_col not in df_attribues.columns:
        raise ValueError(f"Column '{org_col}' missing in df_attribues.")
    
    
    # Hilfsfunktion: Dateinamen-sicherer Org-String
    def _sanitize_filename_part(s: str) -> str:
        s = s.strip()
        # Ungültige Dateisystem-Zeichen ersetzen
        s = re.sub(r'[\\/*?:"<>|]+', "_", s)
        # Windows-reservierte Namen vermeiden
        reserved = {"CON", "PRN", "AUX", "NUL"} | {f"COM{i}" for i in range(1, 10)} | {f"LPT{i}" for i in range(1, 10)}
        if s.upper() in reserved:
            s = f"_{s}_"
        # Nicht-leer
        return s if s else "org"

    os.makedirs(output_dir, exist_ok=True)

    written_files = []

    orgs = df_dataset[org_col].unique()
    for org in orgs:
        org_safe = _sanitize_filename_part(str(org))
        filename_org = f"{org_safe}_{filename}"
        path = os.path.join(output_dir, filename_org)
        df_dataset_org = df_dataset[df_dataset[org_col]==org]
        df_attribues_org = df_attribues[df_attribues[org_col]==org]
        print("write", path)
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df_dataset_org.to_excel(writer, sheet_name="datasets", index=False)
            df_attribues_org.to_excel(writer, sheet_name="attributes", index=False)
        written_files.append(path)

    return written_files

def export_to_excel_sdk_style(df_dataset: pd.DataFrame, df_attribues: pd.DataFrame, 
                              dataset_cols: list, attribute_cols: list, distributions_cols: list,
                              filename="initialimport.xlsx", org_col="author_da_gs", output_dir: str = "."):
    """
    Creates an excel file for every unique organization (e.g. dienstabteilung)
    which contains a sheet for datasets and attribues
    and matches the formats of an SDK Excel Export
    
    :param df_dataset: df with dataset info. Needs to have a org_col
    :type df_dataset: pd.DataFrame
    :param df_attribues: df with attribues info. Needs to have a org_col
    :type df_attribues: pd.DataFrame
    :param filename: End of the filename (gets a prefix from org_col)
    :param org_col: Column name the specifies the dienstabteilung
    """
    
    if org_col not in df_dataset.columns:
        raise ValueError(f"Column '{org_col}' missing in df_dataset.")
    if org_col not in df_attribues.columns:
        raise ValueError(f"Column '{org_col}' missing in df_attribues.")
    
    
    # Hilfsfunktion: Dateinamen-sicherer Org-String
    def _sanitize_filename_part(s: str) -> str:
        s = s.strip()
        # Ungültige Dateisystem-Zeichen ersetzen
        s = re.sub(r'[\\/*?:"<>|]+', "_", s)
        # Windows-reservierte Namen vermeiden
        reserved = {"CON", "PRN", "AUX", "NUL"} | {f"COM{i}" for i in range(1, 10)} | {f"LPT{i}" for i in range(1, 10)}
        if s.upper() in reserved:
            s = f"_{s}_"
        # Nicht-leer
        return s if s else "org"

    os.makedirs(output_dir, exist_ok=True)

    # add empty cols if necessary

    df_dataset = add_empty_cols(df=df_dataset, colnames=dataset_cols)
    df_attribues = add_empty_cols(df=df_attribues, colnames=attribute_cols)
    distributions = add_empty_cols(df=df_dataset, colnames=distributions_cols)

    written_files = []

    orgs = df_dataset[org_col].unique()
    for org in orgs:
        org_safe = _sanitize_filename_part(str(org))
        filename_org = f"{org_safe}_{filename}"
        path = os.path.join(output_dir, filename_org)
        df_dataset_org = df_dataset[df_dataset[org_col]==org]
        df_attribues_org = df_attribues[df_attribues[org_col]==org]
        distributions_org = distributions[distributions[org_col]==org]
        print("write", path)



        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df_dataset_org[dataset_cols].to_excel(writer, sheet_name="Datensätze", index=False, startrow=1)
            df_attribues_org[attribute_cols].to_excel(writer, sheet_name="Bestandteile", index=False, startrow=1)
            distributions_org[distributions_cols].to_excel(writer, sheet_name="Distributionen", index=False, startrow=1)
        
        
            # Alignment einmal definieren (Textumbruch)
            wrap_alignment = Alignment(wrap_text=True)

                    
            # Worksheets holen
            ws_datasets = writer.sheets["Datensätze"]
            ws_parts    = writer.sheets["Bestandteile"]
            ws_dists    = writer.sheets["Distributionen"]

            # Formatierung anwenden (startrow=1 wie oben)
            format_sheet(ws_datasets, ncols=len(dataset_cols),      nrows=len(df_dataset_org),  startrow_excel=1)
            format_sheet(ws_parts,    ncols=len(attribute_cols),    nrows=len(df_attribues_org), startrow_excel=1)
            format_sheet(ws_dists,    ncols=len(distributions_cols), nrows=len(distributions),  startrow_excel=1)


        
        written_files.append(path)

    return written_files

def add_empty_cols(df: pd.DataFrame, colnames: list) -> pd.DatFrame:
    """
    Add empty columns to df if they dont exist
    """
    missing = [c for c in colnames if c not in df.columns]
    for c in missing:
        df[c] = pd.NA  
    return df


# Hilfsfunktion: Formatierung für ein Sheet anwenden
def format_sheet(ws, ncols, nrows, startrow_excel=1, col_width=35, row_height=30):
    """
    ws: openpyxl worksheet
    ncols: Anzahl Spalten (DataFrame-Spalten)
    nrows: Anzahl Datenzeilen (ohne Header)
    startrow_excel: startrow wie bei to_excel (0-basiert), aber hier als int,
                    z.B. startrow=1 bedeutet Header in Excel-Zeile 2
    """

    header_row = startrow_excel + 1      # Excel-Zeile mit Header (bei startrow=1 -> 2)
    first_row = header_row               # wir formatieren ab Header
    last_row = header_row + nrows        # + Datenzeilen

    # 1) Spaltenbreite fix setzen
    for col_idx in range(1, ncols + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    # 2) Zeilenhöhe fix setzen
    for r in range(first_row, last_row + 1):
        ws.row_dimensions[r].height = row_height

    # 3) Textumbruch für alle Zellen im genutzten Bereich setzen
    #    (Header + Datenbereich)
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=1, max_col=ncols):
        for cell in row:
            # bestehende Ausrichtung beibehalten, nur wrap_text setzen:
            cell.alignment = cell.alignment.copy(wrap_text=True)



if __name__ == "__main__":

    df_dataset = pd.DataFrame({"author_da_gs": ["Tiefbauamt (TAZ)", "Verkehrsbetriebe (VBZ)", "Verkehrsbetriebe (VBZ)", "Entsorgung + Recycling Zürich (ERZ)"]})
    export_to_excel_by_org(df_dataset, pd.DataFrame())